from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from thoth.adapters.parsers.common import (
    coverage,
    node_id,
    parser_artifact,
    safe_zip,
    validate_byte_hash,
)
from thoth.domain.artifact import (
    ArtifactEnvelope,
    ParseIssue,
    SourceLocator,
    StructuralDocument,
    StructuralNode,
)
from thoth.domain.enums import ParserErrorCode, StructuralNodeKind
from thoth.domain.errors import ParserFailure


class XlsxParser:
    name = "xlsx"
    version = "1.0.0"
    media_types = frozenset({"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    suffixes = frozenset({".xlsx"})

    def parse(self, artifact: ArtifactEnvelope, raw: bytes) -> StructuralDocument:
        validate_byte_hash(artifact, raw)
        with safe_zip(raw) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise ParserFailure(
                    ParserErrorCode.MEDIA_TYPE_MISMATCH,
                    "XLSX container has no xl/workbook.xml",
                )
            has_external_links = any(
                name.startswith("xl/externalLinks/") for name in archive.namelist()
            )
        try:
            formula_book: Any = load_workbook(
                io.BytesIO(raw), read_only=False, data_only=False, keep_links=False
            )
            value_book: Any = load_workbook(
                io.BytesIO(raw), read_only=False, data_only=True, keep_links=False
            )
        except Exception as exc:
            raise ParserFailure(ParserErrorCode.CORRUPT_DOCUMENT, "invalid XLSX workbook") from exc

        nodes: list[StructuralNode] = []
        warnings: list[ParseIssue] = []
        if has_external_links:
            warnings.append(
                ParseIssue(
                    code=ParserErrorCode.EXTERNAL_LINK_IGNORED,
                    message="external workbook links were not followed",
                )
            )
        for sheet in formula_book.worksheets:
            value_sheet = value_book[sheet.title]
            if sheet.sheet_state != "visible":
                warnings.append(
                    ParseIssue(
                        code=ParserErrorCode.HIDDEN_CONTENT,
                        message=f"sheet {sheet.title!r} is {sheet.sheet_state}",
                        locator=SourceLocator(sheet=sheet.title),
                    )
                )
            merged_ranges = tuple(sheet.merged_cells.ranges)
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    extraction_warnings: list[str] = []
                    if cell.data_type == "f":
                        cached = value_sheet[cell.coordinate].value
                        extraction_warnings.append("FORMULA_NOT_RECALCULATED")
                        text = f"formula={value}; cached={cached!r}"
                    else:
                        text = str(value)
                    if any(cell.coordinate in merged_range for merged_range in merged_ranges):
                        extraction_warnings.append("MERGED_CELL_MEMBER")
                    nodes.append(
                        StructuralNode(
                            node_id=node_id(artifact, len(nodes), "CELL"),
                            artifact_id=artifact.artifact_id,
                            kind=StructuralNodeKind.CELL,
                            ordinal=len(nodes),
                            text=text,
                            locator=SourceLocator(
                                sheet=sheet.title,
                                cell_range=cell.coordinate,
                            ),
                            extraction_warnings=tuple(extraction_warnings),
                        )
                    )
        formula_book.close()
        value_book.close()
        return StructuralDocument(
            artifact=parser_artifact(artifact, name=self.name, version=self.version),
            nodes=tuple(nodes),
            extraction_coverage=coverage(nodes),
            warnings=tuple(warnings),
        )
