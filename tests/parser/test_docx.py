from __future__ import annotations

import io
from typing import Any

from docx import Document

from thoth.adapters.parsers.registry import default_parser_registry
from thoth.domain.enums import StructuralNodeKind


def test_docx_preserves_body_and_table_order(artifact_factory: Any) -> None:
    output = io.BytesIO()
    source = Document()
    source.add_heading("연구 목표", level=1)
    source.add_paragraph("센서 정확도를 검증한다.")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "목표"
    table.cell(0, 1).text = "94%"
    source.save(output)
    raw = output.getvalue()
    artifact = artifact_factory(
        raw,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".docx",
    )

    document = default_parser_registry().parse(artifact, raw)

    kinds = [node.kind for node in document.nodes]
    assert kinds[:3] == [
        StructuralNodeKind.SECTION,
        StructuralNodeKind.PARAGRAPH,
        StructuralNodeKind.TABLE,
    ]
    cells = [node for node in document.nodes if node.kind == StructuralNodeKind.CELL]
    assert [cell.text for cell in cells] == ["목표", "94%"]
    assert cells[1].locator.cell_range == "R1C2"
    assert document.artifact.byte_sha256 == artifact.byte_sha256
